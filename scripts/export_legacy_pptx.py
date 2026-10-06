#!/usr/bin/env python3
"""Exporta un reporte de incidencias en formato JSON (generado por la aplicación)
a una presentación PowerPoint (.pptx) con el diseño y formato antiguo
(idéntico a 2026W39_ReporteIncidencias.pptx).

Uso típico:
    python scripts/export_legacy_pptx.py reporte.json
    python scripts/export_legacy_pptx.py reporte.json -o 2026W39_ReporteIncidencias.pptx
    python scripts/export_legacy_pptx.py reporte.json --template plantilla.pptx
"""

import argparse
from copy import deepcopy
import io
import json
from pathlib import Path
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print(
        "ERROR: falta la librería 'python-pptx'.\n"
        "Instálala con:\n"
        "    pip install python-pptx\n"
        "o (en red corporativa con proxy):\n"
        '    NO_PROXY="*" pip install --trusted-host pypi.org --trusted-host files.pythonhosted.org python-pptx',
        file=sys.stderr,
    )
    sys.exit(1)

ORANGE = RGBColor(255, 120, 0)
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(128, 128, 128)
GREEN = RGBColor(29, 135, 84)
LINE_COLOR = RGBColor(218, 214, 206)
FONT_NAME = "Arial"

LABEL_WORDS = ("Impacto", "Causa", "Solución", "ID", "Fecha", "Duración")
TICKET_RE = re.compile(r"^[A-Z0-9]{6,16}$")

# Layout de slots verticales estándar en diapositiva 16:9 (altura 7.5 pulg.)
SLOT_1 = [0.95]
SLOT_2 = [0.74, 3.90]
SLOT_3 = [0.71, 2.98, 5.04]


def estimate_incident_lines(inc):
    """Estima la cantidad máxima de líneas visuales de una tarjeta para posicionar verticalmente."""
    impact_text = inc.get("impact") or ""
    if inc.get("metrics"):
        impact_text = inc.get("metrics") + "\n" + impact_text
    imp_lines = sum(max(1, len(l) // 45 + (1 if len(l) % 45 else 0)) for l in impact_text.splitlines() if l.strip()) or 1

    cause_text = inc.get("cause") or ""
    cause_lines = sum(max(1, len(l) // 50 + (1 if len(l) % 50 else 0)) for l in cause_text.splitlines() if l.strip()) or 1

    sol_text = inc.get("solution") or ""
    sol_lines = sum(max(1, len(l) // 60 + (1 if len(l) % 60 else 0)) for l in sol_text.splitlines() if l.strip()) or 1
    ap_text = inc.get("actionPoints") or ""
    if ap_text:
        sol_lines += 1
        sol_lines += sum(max(1, len(l) // 60 + (1 if len(l) % 60 else 0)) for l in ap_text.splitlines() if l.strip())

    return max(imp_lines, cause_lines, sol_lines)


def emu_in(v):
    return round(Emu(v).inches, 2) if v is not None else None


def map_group_title(group_name):
    """Mapea el campo `group` de la aplicación al título de la diapositiva del PPT antiguo."""
    g = (group_name or "").strip()
    gl = g.lower()
    if not g:
        return "Incidencias RED"
    if g.startswith("Incidencias ") or g.startswith("Otras "):
        return g
    if "it mm" in gl:
        return "Incidencias IT MM"
    if "it" in gl:
        return "Incidencias IT"
    if "5000" in gl or "5.000" in gl:
        return "Incidencias RED  (> 5000 CLIENTES)"
    if "b2b" in gl:
        return "Incidencias RED  (Impacto en B2B)"
    if "relevantes" in gl or "climatolog" in gl or "escalados" in gl:
        return "Incidencias RED  (Relevantes por duración/Climatología/Escalados RRII)"
    if "otras" in gl:
        if "red (" in gl:
            return "Otras Incidencias RED"
        return "Incidencias RED  (Relevantes por duración/Climatología/Escalados RRII)"
    if gl.startswith("red"):
        return f"Incidencias {g}"
    return f"Incidencias {g}"


def group_sort_priority(group_name):
    """
    Retorna la prioridad de ordenación para los grupos en las diapositivas:
    1. Incidencias IT
    2. Incidencias IT MM (siempre justo después de Incidencias IT)
    3. Incidencias RED (> 5000 CLIENTES)
    4. Incidencias RED (Relevantes por duración/Climatología/Escalados RRII)
    5. Incidencias RED (Impacto en B2B)
    6. Otras Incidencias RED / Otras RED
    7. Resto de grupos
    """
    g = (group_name or "").strip()
    gl = g.lower()
    title = map_group_title(g).lower()

    if "it mm" in gl or "it mm" in title:
        return (2, g)
    if ("it" in gl or "it" in title) and "it mm" not in gl and "it mm" not in title:
        return (1, g)
    if "5000" in gl or "5.000" in gl or "5000" in title:
        return (3, g)
    if "relevantes" in gl or "climatolog" in gl or "escalados" in gl or "relevantes" in title:
        return (4, g)
    if "b2b" in gl or "b2b" in title:
        return (5, g)
    if "otras" in gl or "otras" in title:
        return (6, g)
    return (7, g)


def parse_date_sort_key(date_str):
    """Convierte fechas con formato DD/MM/YYYY HH:MM a clave ordenable YYYY-MM-DD HH:MM."""
    if not date_str:
        return ""
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}:\d{2}))?", str(date_str).strip())
    if m:
        d, mo, y, t = m.group(1).zfill(2), m.group(2).zfill(2), m.group(3), m.group(4) or "00:00"
        return f"{y}-{mo}-{d} {t}"
    return str(date_str)


def set_shape_text_styled(shape, new_text, default_size=10.5, default_color=BLACK, default_bold=False):
    """Asigna texto a un shape preservando o aplicando fuente, tamaño y color corporativo con auto-ajuste de tamaño."""
    tf = shape.text_frame
    font_name = FONT_NAME
    font_bold = default_bold
    font_color = default_color

    text_str = str(new_text or "").strip()
    lines = text_str.split("\n")

    # Si no se pasó un tamaño explícitamente menor, ajustar dinámicamente si el texto es extenso
    if default_size >= 10.5:
        if len(text_str) > 240 or len(lines) >= 6:
            font_size = Pt(9.5)
        elif len(text_str) > 140 or len(lines) >= 4:
            font_size = Pt(10.0)
        else:
            font_size = Pt(default_size)
    else:
        font_size = Pt(default_size)

    if tf.paragraphs and tf.paragraphs[0].runs:
        r0 = tf.paragraphs[0].runs[0]
        if r0.font.name:
            font_name = r0.font.name
        if r0.font.bold is not None:
            font_bold = r0.font.bold
        if r0.font.color and hasattr(r0.font.color, "rgb") and r0.font.color.type == 1:
            font_color = r0.font.color.rgb

    tf.text = ""
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = font_size
        p.font.bold = font_bold
        if font_color:
            p.font.color.rgb = font_color


def set_solution_with_action_points(shape, solution_text, action_points_text):
    """Asigna la solución y, si existen, los Action Points formateados debajo."""
    tf = shape.text_frame
    tf.text = ""

    font_name = FONT_NAME
    base_color = BLACK

    sol_text = str(solution_text or "").strip()
    ap_raw = str(action_points_text or "").strip()

    total_len = len(sol_text) + len(ap_raw)
    sol_size = Pt(9.5) if total_len > 250 else Pt(10.5)
    ap_size = Pt(9.0) if total_len > 250 else Pt(9.5)

    sol_lines = [l for l in sol_text.split("\n") if l.strip()]
    if not sol_lines:
        sol_lines = [""]

    for i, line in enumerate(sol_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = sol_size
        p.font.bold = False
        p.font.color.rgb = base_color

    if ap_raw:
        # Párrafo de encabezado "Action Points:"
        p_head = tf.add_paragraph()
        p_head.space_before = Pt(6)
        r_head = p_head.add_run()
        r_head.text = "Action Points:"
        r_head.font.name = font_name
        r_head.font.size = ap_size
        r_head.font.bold = True
        r_head.font.color.rgb = GREEN

        for line in ap_raw.split("\n"):
            line = line.strip()
            if not line:
                continue

            p_ap = tf.add_paragraph()
            p_ap.font.name = font_name
            p_ap.font.size = ap_size
            p_ap.space_before = Pt(2)

            r_bullet = p_ap.add_run()
            r_bullet.text = "• "
            r_bullet.font.bold = False
            r_bullet.font.color.rgb = base_color

            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 2:
                r_code = p_ap.add_run()
                r_code.text = parts[0]
                r_code.font.bold = True
                r_code.font.color.rgb = base_color

                rest = " | " + " | ".join(parts[1:])
                r_rest = p_ap.add_run()
                r_rest.text = rest
                r_rest.font.bold = False
                r_rest.font.color.rgb = base_color
            else:
                r_text = p_ap.add_run()
                r_text.text = line
                r_text.font.bold = False
                r_text.font.color.rgb = base_color


def populate_group(grp_shape, inc, top_offset_in=None):
    """Rellena una tarjeta de incidencia (GroupShape) clonada de la plantilla."""
    if top_offset_in is not None:
        grp_shape.top = Inches(top_offset_in)

    def collect(shapes, res):
        for s in shapes:
            if s.shape_type == 6:
                collect(s.shapes, res)
            elif s.has_text_frame:
                res.append(s)

    text_shapes = []
    collect(grp_shape.shapes, text_shapes)

    title_val = inc.get("title") or " · ".join(filter(None, [inc.get("category"), inc.get("system")])) or "Incidencia"

    raw_ticket = inc.get("ticket") or ""
    # En el PPT antiguo, los tickets múltiples se formatean uno por línea para que
    # encajen sin solaparse y sean reconocidos por el importador
    ticket_parts = [t.strip() for t in re.split(r"\s*(?://|/)\s*", raw_ticket) if t.strip()]
    ticket_val = "\n".join(ticket_parts) if ticket_parts else raw_ticket

    date_val = inc.get("date") or ""
    dur_val = inc.get("duration") or ""

    # Formateo de métricas / impacto
    impact_lines = []
    if inc.get("metrics"):
        for line in inc["metrics"].strip().split("\n"):
            line = line.strip()
            if line:
                if " | " in line:
                    line = line.replace(" | ", ": ")
                impact_lines.append(line)
    if inc.get("impact") and inc.get("impact").strip() not in impact_lines:
        impact_lines.append(inc["impact"].strip())
    impact_val = "\n".join(impact_lines)

    cause_val = inc.get("cause") or ""
    solution_val = inc.get("solution") or ""
    action_points_val = inc.get("actionPoints") or ""

    for s in text_shapes:
        x = emu_in(s.left)
        y = emu_in(s.top)
        text = s.text_frame.text.strip()
        if not text:
            continue

        lines = [l.strip() for l in text.split("\n") if l.strip()]
        if lines and all(TICKET_RE.fullmatch(l) for l in lines):
            set_shape_text_styled(s, ticket_val, default_size=10.5, default_bold=False)
            continue

        if text in LABEL_WORDS:
            continue

        # Título: x < 1.5 e y < 2.0 (coordenadas relativas canónicas dentro del grupo)
        if x < 1.5 and y < 2.0:
            set_shape_text_styled(s, title_val, default_size=15, default_color=ORANGE, default_bold=False)
            continue

        # Fecha: 9.0 < x < 10.0 e y > 1.4
        if 9.0 < x < 10.0 and y > 1.4:
            set_shape_text_styled(s, date_val, default_size=10.5, default_bold=False)
            continue

        # Duración: x > 10.5 e y > 1.4
        if x > 10.5 and y > 1.4:
            set_shape_text_styled(s, dur_val, default_size=10.5, default_bold=False)
            continue

        # Cuerpo (3 columnas): y > 3.8
        if y > 3.8:
            if abs(x - 0.5) < 1.0:
                set_shape_text_styled(s, impact_val, default_size=10.5, default_bold=False)
            elif abs(x - 4.0) < 1.0:
                set_shape_text_styled(s, cause_val, default_size=10.5, default_bold=False)
            elif abs(x - 8.1) < 1.0:
                set_solution_with_action_points(s, solution_val, action_points_val)


def find_template_pptx(custom_template=None):
    """Localiza el archivo PPTX plantilla con el diseño corporativo antiguo."""
    candidates = []
    if custom_template:
        candidates.append(Path(custom_template))

    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parent

    candidates.extend([
        script_dir / "legacy_template.pptx",
        repo_root / "2026W39_ReporteIncidencias.pptx",
        script_dir / "2026W39_ReporteIncidencias.pptx",
    ])

    for p in candidates:
        if p.is_file():
            return p
    return None


def find_logo_image(custom_logo=None):
    """Localiza el archivo de imagen del logo corporativo de Orange."""
    candidates = []
    if custom_logo:
        candidates.append(Path(custom_logo))

    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parent

    candidates.extend([
        repo_root / "app" / "assets" / "orange-logo.png",
        script_dir / "orange-logo.png",
        repo_root / "app" / "assets" / "orange-logo.svg",
        repo_root / "app" / "assets" / "brands" / "orange.png",
    ])

    for p in candidates:
        if p.is_file():
            return p
    return None


def apply_corporate_logo(prs, logo_path=None):
    """Asegura que el logo corporativo de Orange esté presente en el patrón (Slide Master) y la portada."""
    logo_file = find_logo_image(logo_path)
    if not logo_file:
        return

    # python-pptx requiere un formato raster para add_picture y para manipular el blip
    if logo_file.suffix.lower() == ".svg":
        png_sibling = logo_file.with_suffix(".png")
        if png_sibling.is_file():
            logo_file = png_sibling
        else:
            repo_png = Path(__file__).resolve().parent.parent / "app" / "assets" / "orange-logo.png"
            if repo_png.is_file():
                logo_file = repo_png
            else:
                return

    try:
        with open(logo_file, "rb") as f:
            logo_bytes = f.read()
    except Exception:
        return

    # 1. Actualizar Picture 3 en Slide Master (cabecera superior derecha de diapositivas)
    for master in prs.slide_masters:
        for s in master.shapes:
            if s.name == "Picture 3":
                blip = s._element.xpath(".//a:blip")
                if blip:
                    rId = blip[0].attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rId:
                        try:
                            image_part = master.part.related_part(rId)
                            image_part._blob = logo_bytes
                        except Exception:
                            pass
                size = Inches(0.48)
                s.width = size
                s.height = size
                s.top = Inches(0.09)
                s.left = Inches(12.35)

    # 2. Portada (Slide 0): asegurar que el logo esté presente sobre el título
    if prs.slides:
        slide0 = prs.slides[0]
        cover_logo = next((s for s in slide0.shapes if s.name == "CoverLogo"), None)
        if cover_logo is None:
            try:
                new_pic = slide0.shapes.add_picture(
                    str(logo_file),
                    left=Inches(1.28),
                    top=Inches(2.0),
                    width=Inches(0.95),
                    height=Inches(0.95),
                )
                new_pic.name = "CoverLogo"
            except Exception:
                pass


def sanitize_branding_text(prs):
    """Reemplaza cualquier mención residual de 'MasOrange' por 'Orange' en patrones y diapositivas."""
    for m in prs.slide_masters:
        for s in m.shapes:
            if s.has_text_frame and "masorange" in s.text_frame.text.lower():
                for p in s.text_frame.paragraphs:
                    p.text = re.sub(r"(?i)masorange", "Orange", p.text)
    for slide in prs.slides:
        for s in slide.shapes:
            if s.has_text_frame and "masorange" in s.text_frame.text.lower():
                for p in s.text_frame.paragraphs:
                    p.text = re.sub(r"(?i)masorange", "Orange", p.text)


def clean_template_artifacts(prs):
    """
    Elimina de la plantilla, patrones y layouts cualquier forma o placeholder
    residual que genere marcas de agua indeseadas ("Add tracker", "Source", "subtítulo", "Footnotes").
    """
    keywords = ["tracker", "source", "subtitle", "subtítulo", "footnote"]

    # 1. Limpiar layouts
    for layout in prs.slide_layouts:
        shapes_to_delete = []
        for s in list(layout.shapes):
            name = s.name.lower()
            text = s.text_frame.text.lower() if s.has_text_frame else ""
            if any(kw in name or kw in text for kw in keywords):
                shapes_to_delete.append(s)
        for s in shapes_to_delete:
            try:
                sp = s._element
                sp.getparent().remove(sp)
            except Exception:
                pass

    # 2. Limpiar slide masters
    for master in prs.slide_masters:
        shapes_to_delete = []
        for s in list(master.shapes):
            name = s.name.lower()
            text = s.text_frame.text.lower() if s.has_text_frame else ""
            if any(kw in name or kw in text for kw in keywords):
                shapes_to_delete.append(s)
        for s in shapes_to_delete:
            try:
                sp = s._element
                sp.getparent().remove(sp)
            except Exception:
                pass


def generate_legacy_pptx(report_data_or_path, template_path=None, logo_path=None, cards_per_slide=2):
    """Genera y devuelve la instancia Presentation(pptx) en formato antiguo a partir de un dict o ruta JSON."""
    json_p = None
    if isinstance(report_data_or_path, (dict, list)):
        data = report_data_or_path
    else:
        json_p = Path(report_data_or_path)
        if not json_p.is_file():
            raise FileNotFoundError(f"No existe el archivo JSON: {json_p}")
        with open(json_p, "r", encoding="utf-8") as f:
            data = json.load(f)

    if isinstance(data, list):
        report = {
            "year": 2026,
            "week": 1,
            "range": "",
            "dept": "Customer & Service Operations",
            "incidents": data,
        }
    elif isinstance(data, dict):
        report = data
        if "incidents" not in report and "items" in report:
            report["incidents"] = report["items"]
    else:
        raise ValueError("El JSON debe contener un objeto con 'incidents' o una lista de incidencias.")

    tpl_file = find_template_pptx(template_path)
    if not tpl_file:
        raise FileNotFoundError(
            "No se encontró ninguna plantilla PPTX antigua (buscado 'scripts/legacy_template.pptx' o '2026W39_ReporteIncidencias.pptx').\n"
            "Especifica la ruta con --template / -t."
        )

    prs = Presentation(str(tpl_file))

    # Asegurar el logo corporativo de Orange en la presentación y sanear marcas
    apply_corporate_logo(prs, logo_path)
    sanitize_branding_text(prs)
    clean_template_artifacts(prs)

    # Extraer el elemento XML de la tarjeta de incidencia
    # La plantilla legacy_template tiene Slide 0 (Portada) y Slide 1 (Plantilla de tarjetas)
    # Si se usa un PPT completo (como 2026W39_ReporteIncidencias.pptx), buscar el primer grupo de incidencia
    card_template_elem = None
    slide_index_for_card = 1 if len(prs.slides) > 1 else 0

    slides_list = list(prs.slides)
    for slide in slides_list[slide_index_for_card:]:
        for s in slide.shapes:
            if s.shape_type == 6:  # GROUP
                card_template_elem = deepcopy(s._element)
                break
        if card_template_elem is not None:
            break

    if card_template_elem is None:
        raise ValueError(f"No se encontró ninguna forma de tipo Grupo en la plantilla '{tpl_file}'.")

    # Eliminar del template de tarjeta la línea horizontal inferior fija que corta el texto
    for elem in list(card_template_elem.iter()):
        if elem.tag.endswith("cxnSp"):
            pr = elem.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
            if pr is not None:
                xfrm = pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
                if xfrm is not None:
                    off = xfrm.find("{http://schemas.openxmlformats.org/drawingml/2006/main}off")
                    if off is not None and int(off.get("y", 0)) > 4000000:
                        elem.getparent().remove(elem)

    # Limpiar diapositivas de la plantilla dejando solo la portada (slide 0)
    while len(prs.slides) > 1:
        r_id = prs.slides._sldIdLst[-1].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[-1]

    # Actualizar portada (Slide 0)
    slide0 = prs.slides[0]
    meta = report.get("meta") if isinstance(report.get("meta"), dict) else {}
    year = report.get("year") or meta.get("year", 2026)
    week = report.get("week") if report.get("week") is not None else meta.get("week", "")
    dept = report.get("dept") or meta.get("dept", "Customer & Service Operations")

    for s in slide0.shapes:
        if s.name == "Title":
            s.text = "REPORTE INCIDENCIAS IT + RED"
        elif s.name == "Subtitle":
            s.text = f"{year} - Week {week}".strip(" -")
        elif s.name == "Documenttype":
            s.text = dept

    # Agrupar incidencias según su campo 'group'
    incidents = report.get("incidents", [])
    grouped = {}
    for inc in incidents:
        g = inc.get("group") or "RED"
        grouped.setdefault(g, []).append(inc)

    # Layout de diapositiva de contenido: Layout 2 (Default)
    default_layout = prs.slide_layouts[2] if len(prs.slide_layouts) > 2 else prs.slide_layouts[0]

    total_content_slides = 0
    total_cards = 0

    # Ordenar grupos según prioridad corporativa:
    # 1º IT, 2º IT MM (siempre justo después de Incidencias IT), 3º RED >5000, 4º RED Relevantes, 5º RED B2B, 6º Otras RED
    sorted_group_names = sorted(grouped.keys(), key=group_sort_priority)

    for group_name in sorted_group_names:
        inc_list = grouped[group_name]
        # Ordenar incidencias dentro del grupo cronológicamente por fecha ascendente
        inc_list.sort(key=lambda inc: parse_date_sort_key(inc.get("date", "")))
        slide_title = map_group_title(group_name)

        # Si el grupo tiene exactamente 3 incidencias breves, cabe en 1 diapositiva; si alguna es larga, separar
        if len(inc_list) == 3 and cards_per_slide == 2 and max(estimate_incident_lines(x) for x in inc_list) <= 3:
            chunks = [inc_list]
        else:
            chunks = [inc_list[i:i + cards_per_slide] for i in range(0, len(inc_list), cards_per_slide)]

        for chunk in chunks:
            slide = prs.slides.add_slide(default_layout)
            total_content_slides += 1

            # Eliminar placeholders vacíos de la plantilla para evitar marcas de agua ("Add tracker", "Source", "subtítulo")
            for s in list(slide.shapes):
                if s.is_placeholder and s.placeholder_format.type != 1:
                    try:
                        sp = s._element
                        sp.getparent().remove(sp)
                    except Exception:
                        pass

            # Configurar título de diapositiva
            title_shape = None
            for s in slide.shapes:
                if s.is_placeholder and s.placeholder_format.type == 1:
                    title_shape = s
                    break
            if title_shape:
                title_shape.text = slide_title
                if title_shape.text_frame.paragraphs:
                    p = title_shape.text_frame.paragraphs[0]
                    p.font.name = FONT_NAME
                    p.font.size = Pt(22)
                    p.font.color.rgb = BLACK
            else:
                tb = slide.shapes.add_textbox(Inches(0.61), Inches(0.19), Inches(11.5), Inches(0.5))
                tb.text_frame.text = slide_title
                tb.text_frame.paragraphs[0].font.name = FONT_NAME
                tb.text_frame.paragraphs[0].font.size = Pt(22)
                tb.text_frame.paragraphs[0].font.color.rgb = BLACK

            # Posicionar tarjetas en la diapositiva dinámicamente para evitar solapamientos
            if len(chunk) == 1:
                slots = [0.95]
            elif len(chunk) == 2:
                lines_c1 = estimate_incident_lines(chunk[0])
                c1_bottom = 0.74 + 0.95 + lines_c1 * 0.175
                c2_top = max(3.85, min(4.30, round(c1_bottom + 0.35, 2)))
                slots = [0.74, c2_top]
            elif len(chunk) == 3:
                slots = [0.71, 2.98, 5.04]
            else:
                step = 6.0 / len(chunk)
                slots = [0.85 + i * step for i in range(len(chunk))]

            for idx, inc in enumerate(chunk):
                new_grp = deepcopy(card_template_elem)
                slide.shapes._spTree.append(new_grp)
                grp_shape = slide.shapes[-1]
                populate_group(grp_shape, inc, top_offset_in=slots[idx])
                total_cards += 1

                # Si es la primera tarjeta de un par, añadir separador horizontal limpio entre tarjetas
                if len(chunk) == 2 and idx == 0:
                    sep_y = slots[1] - 0.16
                    line_shape = slide.shapes.add_shape(
                        1,  # MSO_SHAPE.RECTANGLE
                        Inches(0.48), Inches(sep_y), Inches(12.24), Pt(1)
                    )
                    line_shape.fill.solid()
                    line_shape.fill.fore_color.rgb = ORANGE
                    line_shape.line.color.rgb = ORANGE
                    line_shape.line.width = Pt(1)

    return prs, total_content_slides, total_cards, tpl_file, year, week, json_p


def generate_legacy_pptx_bytes(report_data_or_path, template_path=None, logo_path=None, cards_per_slide=2) -> bytes:
    """Genera la presentación PPTX legacy y devuelve su contenido en bytes en memoria."""
    prs, _, _, _, _, _, _ = generate_legacy_pptx(
        report_data_or_path, template_path=template_path, logo_path=logo_path, cards_per_slide=cards_per_slide
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_legacy_pptx(json_path, output_path=None, template_path=None, logo_path=None, cards_per_slide=2, quiet=False):
    """Lee el JSON del reporte y genera la presentación PPTX en formato antiguo guardándola en archivo."""
    prs, total_content_slides, total_cards, tpl_file, year, week, json_p = generate_legacy_pptx(
        json_path, template_path=template_path, logo_path=logo_path, cards_per_slide=cards_per_slide
    )

    # Determinar ruta de salida
    if not output_path:
        stem = json_p.stem if json_p else "ReporteIncidencias"
        if stem.endswith("_ReporteIncidencias"):
            stem = stem[:-len("_ReporteIncidencias")]
        w_part = f"{year}W{str(week).zfill(2)}" if week else stem
        parent_dir = json_p.parent if json_p else Path.cwd()
        output_path = parent_dir / f"{w_part}_ReporteIncidencias.pptx"
    else:
        output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(output_path))
    except PermissionError:
        alt_path = output_path.with_name(f"{output_path.stem}_actualizado{output_path.suffix}")
        try:
            prs.save(str(alt_path))
            print(f"\n⚠️  AVISO: '{output_path.name}' está abierto en PowerPoint y bloqueado contra escritura.", file=sys.stderr)
            print(f"    Se ha guardado la versión actualizada en: {alt_path.name}\n", file=sys.stderr)
            output_path = alt_path
        except Exception:
            raise PermissionError(
                f"No se pudo guardar '{output_path.name}' porque está abierto en otra aplicación (PowerPoint).\n"
                f"Ciérralo y vuelve a intentarlo o especifica otra ruta con -o."
            )

    if not quiet:
        print(f"Generado con éxito: {output_path}")
        print(f"  - Total diapositivas: {len(prs.slides)} (1 portada + {total_content_slides} de contenido)")
        print(f"  - Incidencias exportadas: {total_cards}")
        print(f"  - Plantilla base: {tpl_file.name}")

    return output_path


def main():
    parser = argparse.ArgumentParser(
        description="Genera una presentación PPTX en formato antiguo (legacy) a partir de un JSON de incidencias.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ejemplos:
  %(prog)s 2026W39_ReporteIncidencias.json
  %(prog)s reporte.json -o salida.pptx
  %(prog)s reporte.json --template scripts/legacy_template.pptx
        """,
    )
    parser.add_argument("input", nargs="?", help="Ruta al archivo JSON de entrada.")
    parser.add_argument("-i", "--input-file", dest="input_opt", help="Alternativa para especificar el JSON de entrada.")
    parser.add_argument("-o", "--output", help="Ruta del archivo PPTX generado (por defecto AAAAWNN_ReporteIncidencias.pptx).")
    parser.add_argument("-t", "--template", help="Ruta a una plantilla PPTX personalizada con el formato antiguo.")
    parser.add_argument("-l", "--logo", help="Ruta al archivo del logo corporativo (SVG o PNG). Por defecto app/assets/orange-logo.png.")
    parser.add_argument("--cards-per-slide", type=int, default=2, help="Máximo número de tarjetas por diapositiva (por defecto 2).")
    parser.add_argument("-q", "--quiet", action="store_true", help="Suprime mensajes informativos.")

    args = parser.parse_args()
    input_file = args.input or args.input_opt
    if not input_file:
        parser.print_help()
        sys.exit(1)

    try:
        export_legacy_pptx(
            json_path=input_file,
            output_path=args.output,
            template_path=args.template,
            logo_path=args.logo,
            cards_per_slide=args.cards_per_slide,
            quiet=args.quiet,
        )
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
